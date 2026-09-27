"""Declaracion para SINADER del Modulo 9, Parte A (CU-67).

Agrupa por cliente generador el material recibido en el periodo y produce
la planilla XLSX que el administrador carga a mano en la plataforma de la
autoridad. El sistema no se conecta con SINADER.

Supuesto documentado: los datos obligatorios del generador para declarar
son su RUT y su direccion (ademas de la razon social, que siempre existe).
Si la autoridad exige otro campo, basta con agregarlo a CAMPOS_OBLIGATORIOS.
"""
from collections import OrderedDict
from decimal import Decimal
from io import BytesIO

from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from common.trazas import traza
from recepcion.models import Recepcion

from .models import DeclaracionSinader

CERO = Decimal("0.00")
CAMPOS_OBLIGATORIOS = (("rut", "RUT"), ("direccion", "direccion"))


class DeclaracionNoGenerable(Exception):
    def __init__(self, motivo, datos=None):
        super().__init__(motivo)
        self.motivo = motivo
        self.datos = datos or {}


def faltantes_del_cliente(cliente):
    return [etiqueta for campo, etiqueta in CAMPOS_OBLIGATORIOS if not getattr(cliente, campo)]


def compilar_periodo(periodo_inicio, periodo_fin):
    """Separa las descargas del periodo en declarables y excluidas (CU-67).

    Se excluyen las descargas de clientes sin datos obligatorios, y se
    informa que falta completar (Excepcion 2). Sin descargas recibidas en el
    periodo no hay nada que declarar (Excepcion 1).
    """
    recepciones = (
        Recepcion.objects.filter(
            estado=Recepcion.RECIBIDA,
            fecha__gte=periodo_inicio,
            fecha__lte=periodo_fin,
        )
        .select_related("cliente")
        .prefetch_related("detalles__material")
        .order_by("cliente__razon_social", "fecha", "id")
    )
    if not recepciones.exists():
        raise DeclaracionNoGenerable(
            "No hay descargas recibidas en el periodo; no hay material que declarar."
        )

    filas = []
    clientes = OrderedDict()
    excluidos = OrderedDict()
    for recepcion in recepciones:
        cliente = recepcion.cliente
        faltantes = faltantes_del_cliente(cliente)
        if faltantes:
            excl = excluidos.setdefault(
                cliente.pk,
                {"cliente": cliente.razon_social, "faltantes": faltantes, "descargas": 0},
            )
            excl["descargas"] += 1
            continue
        resumen = clientes.setdefault(
            cliente.pk,
            {
                "cliente": cliente.razon_social,
                "rut": cliente.rut,
                "direccion": cliente.direccion,
                "descargas": 0,
                "volumen_m3": CERO,
                "peso_kg": CERO,
            },
        )
        resumen["descargas"] += 1
        for detalle in recepcion.detalles.all():
            peso = detalle.peso_derivado_kg or CERO
            filas.append(
                {
                    "cliente": cliente.razon_social,
                    "rut": cliente.rut,
                    "direccion": cliente.direccion,
                    "fecha": recepcion.fecha,
                    "recepcion": recepcion.pk,
                    "material": detalle.material.nombre,
                    "volumen_m3": detalle.volumen_m3,
                    "peso_kg": peso,
                }
            )
            resumen["volumen_m3"] += detalle.volumen_m3
            resumen["peso_kg"] += peso

    if not filas:
        raise DeclaracionNoGenerable(
            "Ninguna descarga del periodo se puede declarar: todos los clientes "
            "tienen datos obligatorios incompletos.",
            {"excluidos": list(excluidos.values())},
        )
    return filas, list(clientes.values()), list(excluidos.values())


def _armar_planilla(periodo_inicio, periodo_fin, filas, clientes):
    """Planilla XLSX con una hoja de detalle y otra de resumen por cliente."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Declaracion"
    encabezado = [
        "Cliente generador", "RUT", "Direccion", "Fecha", "Recepcion",
        "Material", "Volumen (m3)", "Peso (kg)",
    ]
    hoja.append(encabezado)
    for fila in filas:
        hoja.append(
            [
                fila["cliente"], fila["rut"], fila["direccion"], fila["fecha"],
                fila["recepcion"], fila["material"],
                float(fila["volumen_m3"]), float(fila["peso_kg"]),
            ]
        )
    resumen = libro.create_sheet("Resumen por cliente")
    resumen.append(["Periodo", f"{periodo_inicio.isoformat()} a {periodo_fin.isoformat()}"])
    resumen.append([])
    resumen.append(["Cliente generador", "RUT", "Descargas", "Volumen (m3)", "Peso (kg)"])
    for c in clientes:
        resumen.append(
            [c["cliente"], c["rut"], c["descargas"], float(c["volumen_m3"]), float(c["peso_kg"])]
        )
    for ws, ancho in ((hoja, 18), (resumen, 22)):
        for celda in ws[1] if ws is hoja else ws[3]:
            celda.font = Font(bold=True)
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
    for celda in hoja["D"][1:]:
        celda.number_format = "YYYY-MM-DD"
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def generar_declaracion(periodo_inicio, periodo_fin, usuario):
    """Genera y persiste la declaracion del periodo con su planilla (CU-67)."""
    filas, clientes, excluidos = compilar_periodo(periodo_inicio, periodo_fin)
    total_volumen = sum((f["volumen_m3"] for f in filas), CERO)
    total_peso = sum((f["peso_kg"] for f in filas), CERO)
    contenido = {
        "periodo": {"inicio": periodo_inicio.isoformat(), "fin": periodo_fin.isoformat()},
        "clientes": [
            {**c, "volumen_m3": f"{c['volumen_m3']:.2f}", "peso_kg": f"{c['peso_kg']:.2f}"}
            for c in clientes
        ],
        "excluidos": excluidos,
        "totales": {
            "clientes": len(clientes),
            "descargas": sum(c["descargas"] for c in clientes),
            "volumen_m3": f"{total_volumen:.2f}",
            "peso_kg": f"{total_peso:.2f}",
        },
    }
    planilla = _armar_planilla(periodo_inicio, periodo_fin, filas, clientes)
    with transaction.atomic():
        declaracion = DeclaracionSinader(
            periodo_inicio=periodo_inicio,
            periodo_fin=periodo_fin,
            usuario=usuario if usuario and usuario.is_authenticated else None,
            contenido=contenido,
        )
        nombre = f"sinader-{periodo_inicio.isoformat()}-{periodo_fin.isoformat()}.xlsx"
        declaracion.archivo.save(nombre, ContentFile(planilla), save=True)
    traza(
        "CU-67",
        "declaracion.generada",
        desde=periodo_inicio,
        hasta=periodo_fin,
        clientes=len(clientes),
        excluidos=len(excluidos),
        peso_kg=contenido["totales"]["peso_kg"],
    )
    return declaracion
