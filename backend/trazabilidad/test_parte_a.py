"""Tests del Modulo 9, Parte A: certificados y declaracion SINADER.

Cubren los criterios de aceptacion 1 a 3 de la spec M9 (el criterio 4, la
trazabilidad de la venta, vive en comercial.tests) y las excepciones de los
CU-65 a CU-67.
"""
import tempfile
from datetime import date, time
from decimal import Decimal
from io import BytesIO

from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

# Las planillas generadas en los tests van a un directorio temporal, no al
# media real del proyecto.
MEDIA_TEMPORAL = tempfile.mkdtemp(prefix="triverde-test-media-")

from acceso.models import BitacoraAuditoria, Rol, Usuario
from mantenedores.models import Cliente, Material, Transportista
from recepcion.models import DetalleRecepcion, Recepcion

from .models import CertificadoTrazabilidad, DeclaracionSinader


@override_settings(MEDIA_ROOT=MEDIA_TEMPORAL)
class ParteATestBase(APITestCase):
    @classmethod
    def setUpTestData(cls):
        cls.rol_admin = Rol.objects.create(nombre=Rol.ADMINISTRADOR)
        cls.rol_operador = Rol.objects.create(nombre=Rol.OPERADOR)
        cls.admin = Usuario.objects.create_user(
            username="admin-m9a", password="clave-segura", rol=cls.rol_admin,
            nombre_completo="Admin M9A",
        )
        cls.operador = Usuario.objects.create_user(
            username="operador-m9a", password="clave-segura", rol=cls.rol_operador,
            nombre_completo="Operador M9A",
        )
        cls.cliente = Cliente.objects.create(
            razon_social="Forestal Calle Calle", rut="76.111.222-3",
            direccion="Camino a Quilapilun km 4",
        )
        cls.cliente_incompleto = Cliente.objects.create(
            razon_social="Podas Sin Datos",
        )
        cls.transportista = Transportista.objects.create(
            nombre="Transportes Colina", cliente=cls.cliente
        )
        cls.rama = Material.objects.create(
            nombre="Ramas de poda", categoria=Material.SECA,
            densidad_kg_m3=Decimal("220.00"), admite_chip=True,
            factor_reduccion_chip=Decimal("3.00"),
        )
        cls.pasto = Material.objects.create(
            nombre="Hojas y cesped", categoria=Material.VERDE,
            densidad_kg_m3=Decimal("350.00"),
        )

    def setUp(self):
        self.client.force_authenticate(user=self.admin)

    def crear_recepcion(self, cliente=None, estado=Recepcion.RECIBIDA, fecha=None):
        return Recepcion.objects.create(
            cliente=cliente or self.cliente,
            transportista=self.transportista,
            conductor="Juan Perez",
            fecha=fecha or date(2026, 8, 12),
            hora=time(10, 30),
            estado=estado,
        )

    def agregar_detalle(self, recepcion, material=None, volumen="10.00", peso="2200.00"):
        return DetalleRecepcion.objects.create(
            recepcion=recepcion,
            material=material or self.rama,
            volumen_m3=Decimal(volumen),
            peso_derivado_kg=Decimal(peso) if peso is not None else None,
        )

    def descarga_certificable(self, **kwargs):
        recepcion = self.crear_recepcion(**kwargs)
        self.agregar_detalle(recepcion)
        return recepcion


# --- CU-65: certificado por descarga --------------------------------------


class CertificadoDescargaTests(ParteATestBase):
    """Criterio de aceptacion 1."""

    def _generar(self, recepcion_id):
        return self.client.post(
            reverse("certificado-generar-descarga"), {"recepcion": recepcion_id}, format="json"
        )

    def test_descarga_recibida_con_peso_emite_certificado(self):
        recepcion = self.descarga_certificable()
        self.agregar_detalle(recepcion, material=self.pasto, volumen="4.00", peso="1400.00")
        respuesta = self._generar(recepcion.pk)
        self.assertEqual(respuesta.status_code, status.HTTP_201_CREATED)
        self.assertTrue(respuesta.data["codigo"].startswith("CTR-"))
        self.assertEqual(respuesta.data["tipo"], "descarga")
        contenido = respuesta.data["contenido"]
        self.assertEqual(contenido["cliente"]["razon_social"], "Forestal Calle Calle")
        self.assertEqual(contenido["transportista"], "Transportes Colina")
        self.assertEqual(contenido["totales"], {"volumen_m3": "14.00", "peso_kg": "3600.00"})
        self.assertEqual(len(contenido["materiales"]), 2)
        self.assertEqual(contenido["codigo"], respuesta.data["codigo"])
        self.assertTrue(
            BitacoraAuditoria.objects.filter(
                accion="Generacion de certificado de trazabilidad"
            ).exists()
        )

    def test_descarga_en_curso_no_se_certifica(self):
        """CU-65, Excepcion 1."""
        recepcion = self.descarga_certificable(estado=Recepcion.EN_CURSO)
        respuesta = self._generar(recepcion.pk)
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("recibida", respuesta.data["detalle"])
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 0)

    def test_descarga_rechazada_no_se_certifica(self):
        recepcion = self.descarga_certificable(estado=Recepcion.RECHAZADA)
        self.assertEqual(self._generar(recepcion.pk).status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 0)

    def test_linea_sin_peso_no_se_certifica(self):
        """CU-65, Excepcion 2: el peso derivado no admite nulos; "sin peso" es cero."""
        recepcion = self.crear_recepcion()
        self.agregar_detalle(recepcion, peso="0.00")
        respuesta = self._generar(recepcion.pk)
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(respuesta.data["sin_peso"], ["Ramas de poda"])
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 0)

    def test_la_misma_descarga_no_duplica_el_folio(self):
        recepcion = self.descarga_certificable()
        primero = self._generar(recepcion.pk)
        segundo = self._generar(recepcion.pk)
        self.assertEqual(segundo.status_code, status.HTTP_200_OK)
        self.assertEqual(segundo.data["codigo"], primero.data["codigo"])
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 1)

    def test_folios_correlativos_por_anio(self):
        a = self._generar(self.descarga_certificable().pk).data["codigo"]
        b = self._generar(self.descarga_certificable().pk).data["codigo"]
        anio = date.today().year
        self.assertEqual(a, f"CTR-{anio}-0001")
        self.assertEqual(b, f"CTR-{anio}-0002")

    def test_exportar_devuelve_el_contenido(self):
        recepcion = self.descarga_certificable()
        certificado_id = self._generar(recepcion.pk).data["id"]
        respuesta = self.client.get(reverse("certificado-exportar", args=[certificado_id]))
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        self.assertEqual(respuesta.data["recepcion"], recepcion.pk)
        self.assertIn("codigo", respuesta.data)

    def test_recepcion_inexistente(self):
        self.assertEqual(self._generar(99999).status_code, status.HTTP_404_NOT_FOUND)


# --- CU-66: consolidado mensual ------------------------------------------


class ConsolidadoMensualTests(ParteATestBase):
    """Criterio de aceptacion 2."""

    PERIODO = {"periodo_inicio": "2026-08-01", "periodo_fin": "2026-08-31"}

    def _generar(self, cliente=None, confirmar=False, **periodo):
        datos = {"cliente": (cliente or self.cliente).pk, **(periodo or self.PERIODO)}
        if confirmar:
            datos["confirmar"] = True
        return self.client.post(reverse("certificado-generar-consolidado"), datos, format="json")

    def test_agrupa_por_material_sumando_volumen_y_peso(self):
        r1 = self.crear_recepcion(fecha=date(2026, 8, 3))
        self.agregar_detalle(r1, self.rama, "10.00", "2200.00")
        self.agregar_detalle(r1, self.pasto, "5.00", "1750.00")
        r2 = self.crear_recepcion(fecha=date(2026, 8, 20))
        self.agregar_detalle(r2, self.rama, "6.00", "1320.00")
        # Fuera del periodo y de otro cliente: no entran.
        self.agregar_detalle(self.crear_recepcion(fecha=date(2026, 7, 30)))
        self.agregar_detalle(self.crear_recepcion(cliente=self.cliente_incompleto))
        # En curso: no cuenta aunque este en el periodo.
        self.agregar_detalle(self.crear_recepcion(estado=Recepcion.EN_CURSO))

        respuesta = self._generar()
        self.assertEqual(respuesta.status_code, status.HTTP_201_CREATED)
        contenido = respuesta.data["contenido"]
        materiales = {m["material"]: m for m in contenido["materiales"]}
        self.assertEqual(materiales["Ramas de poda"]["volumen_m3"], "16.00")
        self.assertEqual(materiales["Ramas de poda"]["peso_kg"], "3520.00")
        self.assertEqual(materiales["Ramas de poda"]["descargas"], 2)
        self.assertEqual(materiales["Hojas y cesped"]["volumen_m3"], "5.00")
        self.assertEqual(contenido["totales"]["descargas"], 2)
        self.assertEqual(contenido["totales"]["peso_kg"], "5270.00")
        self.assertEqual(contenido["version"], 1)
        self.assertEqual(respuesta.data["periodo_inicio"], "2026-08-01")
        self.assertIsNone(respuesta.data["recepcion"])

    def test_sin_descargas_en_el_periodo_no_genera(self):
        """CU-66, Excepcion 1."""
        respuesta = self._generar()
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("no tiene descargas", respuesta.data["detalle"])
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 0)

    def test_consolidado_previo_exige_confirmacion_y_conserva_versiones(self):
        """CU-66, Excepcion 2."""
        self.agregar_detalle(self.crear_recepcion())
        primero = self._generar()
        self.assertEqual(primero.status_code, status.HTTP_201_CREATED)

        repetido = self._generar()
        self.assertEqual(repetido.status_code, status.HTTP_409_CONFLICT)
        self.assertTrue(repetido.data["requiere_confirmacion"])
        self.assertEqual(repetido.data["previo"]["codigo"], primero.data["codigo"])
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 1)

        confirmado = self._generar(confirmar=True)
        self.assertEqual(confirmado.status_code, status.HTTP_201_CREATED)
        self.assertEqual(confirmado.data["contenido"]["version"], 2)
        self.assertNotEqual(confirmado.data["codigo"], primero.data["codigo"])
        self.assertEqual(CertificadoTrazabilidad.objects.count(), 2)

    def test_periodo_invalido(self):
        self.agregar_detalle(self.crear_recepcion())
        r = self._generar(periodo_inicio="2026-08-31", periodo_fin="2026-08-01")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)
        r = self._generar(periodo_inicio="agosto", periodo_fin="2026-08-31")
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)

    def test_historial_filtra_por_cliente_y_tipo(self):
        self.agregar_detalle(self.crear_recepcion())
        self._generar()
        recepcion = self.descarga_certificable(cliente=self.cliente_incompleto)
        self.client.post(
            reverse("certificado-generar-descarga"), {"recepcion": recepcion.pk}, format="json"
        )
        todos = self.client.get(reverse("certificado-list"))
        self.assertEqual(len(todos.data), 2)
        del_cliente = self.client.get(reverse("certificado-list"), {"cliente": self.cliente.pk})
        self.assertEqual(len(del_cliente.data), 1)
        consolidados = self.client.get(reverse("certificado-list"), {"tipo": "consolidado"})
        self.assertEqual(len(consolidados.data), 1)
        self.assertEqual(consolidados.data[0]["referencia"], "2026-08-01 a 2026-08-31")


# --- CU-67: declaracion SINADER --------------------------------------------


class DeclaracionSinaderTests(ParteATestBase):
    """Criterio de aceptacion 3."""

    PERIODO = {"periodo_inicio": "2026-08-01", "periodo_fin": "2026-08-31"}

    def _generar(self, **periodo):
        return self.client.post(
            reverse("declaracion-sinader-generar"), periodo or self.PERIODO, format="json"
        )

    def test_agrupa_por_cliente_y_excluye_sin_datos_obligatorios(self):
        r1 = self.crear_recepcion(fecha=date(2026, 8, 5))
        self.agregar_detalle(r1, self.rama, "10.00", "2200.00")
        r2 = self.crear_recepcion(fecha=date(2026, 8, 9))
        self.agregar_detalle(r2, self.pasto, "4.00", "1400.00")
        self.agregar_detalle(self.crear_recepcion(cliente=self.cliente_incompleto))

        respuesta = self._generar()
        self.assertEqual(respuesta.status_code, status.HTTP_201_CREATED)
        contenido = respuesta.data["contenido"]
        self.assertEqual(len(contenido["clientes"]), 1)
        incluido = contenido["clientes"][0]
        self.assertEqual(incluido["cliente"], "Forestal Calle Calle")
        self.assertEqual(incluido["descargas"], 2)
        self.assertEqual(incluido["peso_kg"], "3600.00")
        self.assertEqual(
            contenido["excluidos"],
            [{"cliente": "Podas Sin Datos", "faltantes": ["RUT", "direccion"], "descargas": 1}],
        )
        self.assertEqual(contenido["totales"]["peso_kg"], "3600.00")
        self.assertTrue(respuesta.data["nombre_archivo"].endswith(".xlsx"))
        self.assertTrue(
            BitacoraAuditoria.objects.filter(accion="Generacion de declaracion SINADER").exists()
        )

    def test_la_planilla_tiene_el_detalle_y_el_resumen(self):
        r1 = self.crear_recepcion(fecha=date(2026, 8, 5))
        self.agregar_detalle(r1, self.rama, "10.00", "2200.00")
        self.agregar_detalle(r1, self.pasto, "4.00", "1400.00")
        declaracion_id = self._generar().data["id"]

        respuesta = self.client.get(reverse("declaracion-sinader-descargar", args=[declaracion_id]))
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml", respuesta["Content-Type"])
        libro = load_workbook(BytesIO(b"".join(respuesta.streaming_content)))
        detalle = libro["Declaracion"]
        filas = list(detalle.iter_rows(values_only=True))
        self.assertEqual(filas[0][0], "Cliente generador")
        self.assertEqual(len(filas), 3)
        self.assertEqual(filas[1][1], "76.111.222-3")
        self.assertEqual(filas[1][5], "Ramas de poda")
        self.assertEqual(filas[1][7], 2200.0)
        resumen = libro["Resumen por cliente"]
        fila_cliente = list(resumen.iter_rows(min_row=4, values_only=True))[0]
        self.assertEqual(fila_cliente[0], "Forestal Calle Calle")
        self.assertEqual(fila_cliente[2], 1)
        self.assertEqual(fila_cliente[4], 3600.0)

    def test_periodo_sin_descargas_no_genera(self):
        """CU-67, Excepcion 1."""
        respuesta = self._generar()
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(DeclaracionSinader.objects.count(), 0)

    def test_todos_excluidos_no_genera_y_avisa(self):
        """CU-67, Excepcion 2 en su forma extrema."""
        self.agregar_detalle(self.crear_recepcion(cliente=self.cliente_incompleto))
        respuesta = self._generar()
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(respuesta.data["excluidos"][0]["faltantes"], ["RUT", "direccion"])
        self.assertEqual(DeclaracionSinader.objects.count(), 0)

    def test_historial_lista_las_declaraciones(self):
        self.agregar_detalle(self.crear_recepcion())
        self._generar()
        respuesta = self.client.get(reverse("declaracion-sinader-list"))
        self.assertEqual(len(respuesta.data), 1)
        self.assertEqual(respuesta.data[0]["usuario_nombre"], "Admin M9A")


# --- Permisos --------------------------------------------------------------


class PermisosParteATests(ParteATestBase):
    def test_el_operador_no_accede(self):
        self.client.force_authenticate(user=self.operador)
        recepcion = self.descarga_certificable()
        casos = [
            self.client.get(reverse("certificado-list")),
            self.client.post(
                reverse("certificado-generar-descarga"), {"recepcion": recepcion.pk}, format="json"
            ),
            self.client.post(
                reverse("certificado-generar-consolidado"),
                {"cliente": self.cliente.pk, **ConsolidadoMensualTests.PERIODO},
                format="json",
            ),
            self.client.get(reverse("declaracion-sinader-list")),
            self.client.post(
                reverse("declaracion-sinader-generar"), DeclaracionSinaderTests.PERIODO, format="json"
            ),
        ]
        for respuesta in casos:
            self.assertEqual(respuesta.status_code, status.HTTP_403_FORBIDDEN)

    def test_sin_token_no_accede(self):
        self.client.force_authenticate(user=None)
        self.assertEqual(
            self.client.get(reverse("certificado-list")).status_code, status.HTTP_401_UNAUTHORIZED
        )
